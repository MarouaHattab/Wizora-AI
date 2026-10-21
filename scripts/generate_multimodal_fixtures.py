"""Generate inspectable lesson fixtures, including genuine synthesized speech.

Windows uses its installed offline System.Speech voice. No ASR or LLM output is
fabricated. Pass an existing WAV on other systems with --speech-wav.
"""
import argparse
import os
from pathlib import Path
import shutil
import subprocess

LESSON = (
    "Solar learning laboratory. Solar panels convert sunlight into electrical energy. "
    "A battery stores energy for use after sunset. An inverter converts direct current into alternating current. "
    "The laboratory has six solar panels. Each panel produces two hundred watts in the test conditions. "
    "Together the six panels produce twelve hundred watts. The battery stores three kilowatt hours of energy. "
    "The laboratory measures power in watts and stored energy in kilowatt hours. "
    "Cloud cover can reduce solar power production. The experiment compares sunny and cloudy conditions. "
    "Students record power every ten minutes. The experiment does not measure installation costs."
)
VISUAL = [
    "SOLAR LEARNING LAB\n\n6 solar panels x 200 watts = 1200 watts\nBattery storage: 3 kilowatt hours\nInverter: direct current to alternating current",
    "VISUAL-ONLY LESSON\n\nThe experiment access code is VIOLET 73.\nMeasure power every 10 minutes.\nInstallation costs were NOT measured.",
]


def generate(destination, speech_wav=None):
    import fitz
    import imageio_ffmpeg
    from docx import Document
    from PIL import Image, ImageDraw, ImageFont
    destination = Path(destination).resolve()
    destination.mkdir(parents=True, exist_ok=True)
    (destination / "lesson.txt").write_text(LESSON, encoding="utf-8")
    long_lesson = "\n\n".join(f"Experiment record {number}. {LESSON}" for number in range(1, 25))
    (destination / "lesson-long.txt").write_text(long_lesson + "\n\nFINAL CHECKPOINT: The reserve identifier is ORCHID 92.", encoding="utf-8")
    (destination / "lesson-utf16.txt").write_text("Énergie solaire\n\n" + LESSON, encoding="utf-16")
    pdf = fitz.open()
    for text in (LESSON[:LESSON.index("The laboratory measures")], LESSON[LESSON.index("The laboratory measures"):]):
        page = pdf.new_page()
        page.insert_textbox(fitz.Rect(54, 54, 540, 760), text, fontsize=14)
    pdf.save(destination / "lesson.pdf")
    pdf.close()
    word = Document()
    word.add_heading("Solar learning laboratory", 0)
    word.add_paragraph(LESSON)
    word.add_heading("Experimental values", 1)
    table = word.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Panels", "6"
    row = table.add_row()
    row.cells[0].text, row.cells[1].text = "Power per panel", "200 watts"
    word.add_paragraph("The experiment does not measure installation costs.")
    word.save(destination / "lesson.docx")
    wav = destination / "lesson.wav"
    if speech_wav:
        shutil.copyfile(speech_wav, wav)
    else:
        environment = dict(os.environ, WIZORA_FIXTURE_TEXT=str(destination / "lesson.txt"), WIZORA_FIXTURE_WAV=str(wav))
        command = (
            "$ErrorActionPreference='Stop'; Add-Type -AssemblyName System.Speech; "
            "$voice = New-Object System.Speech.Synthesis.SpeechSynthesizer; "
            "$voice.SelectVoice('Microsoft Zira Desktop'); "
            "$voice.SetOutputToWaveFile($env:WIZORA_FIXTURE_WAV); "
            "$voice.Speak([IO.File]::ReadAllText($env:WIZORA_FIXTURE_TEXT)); $voice.Dispose()"
        )
        subprocess.run(["powershell.exe", "-NoProfile", "-NonInteractive", "-Command", command],
                       env=environment, check=True, capture_output=True, timeout=60)
    ffmpeg = imageio_ffmpeg.get_ffmpeg_exe()
    def convert(*arguments):
        subprocess.run([ffmpeg, "-v", "error", "-nostdin", "-y", *map(str, arguments)],
                       check=True, capture_output=True, timeout=120)
    for extension, codec in (("mp3", "libmp3lame"), ("m4a", "aac")):
        convert("-i", wav, "-c:a", codec, destination / ("lesson." + extension))
    font_path = Path("C:/Windows/Fonts/arial.ttf")
    font = ImageFont.truetype(str(font_path), 40) if font_path.exists() else ImageFont.load_default(size=36)
    for number, text in enumerate(VISUAL):
        image = Image.new("RGB", (1280, 720), "#f7f5ff")
        draw = ImageDraw.Draw(image)
        draw.multiline_text((70, 90), text, fill="#312e81", font=font, spacing=18)
        image.save(destination / f"slide-{number}.png")
    for extension in ("mp4", "mov", "webm"):
        codec = "libvpx-vp9" if extension == "webm" else "libx264"
        audio_codec = "libopus" if extension == "webm" else "aac"
        convert("-loop", "1", "-framerate", "1", "-i", destination / "slide-0.png", "-i", wav,
                "-c:v", codec, "-pix_fmt", "yuv420p", "-c:a", audio_codec, "-shortest", "-threads", "2",
                destination / ("lesson." + extension))
    # No audio stream: the access code can only be recovered from visible pixels.
    convert("-loop", "1", "-framerate", "1", "-i", destination / "slide-1.png", "-t", "4", "-an",
            "-c:v", "libx264", "-pix_fmt", "yuv420p", destination / "visual-only.mp4")
    (destination / "empty.txt").write_bytes(b"")
    (destination / "corrupt.pdf").write_bytes(b"%PDF-this is not a valid PDF")
    (destination / "corrupt.docx").write_bytes(b"PK corrupt archive")
    (destination / "invalid.wav").write_bytes(b"this is not audio")
    (destination / "unsupported.exe").write_bytes(b"unsupported")
    print(f"Generated fixtures in {destination}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("destination", nargs="?", default="output/multimodal-fixtures")
    parser.add_argument("--speech-wav")
    args = parser.parse_args()
    generate(args.destination, args.speech_wav)
