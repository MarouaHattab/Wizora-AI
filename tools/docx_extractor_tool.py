import hashlib
import json
import re
from pathlib import Path
from zipfile import BadZipFile, ZipFile

from docx import Document as WordDocument
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph
from langchain_core.documents import Document
from lxml import etree

from helper.ingestion_errors import (
    AssetNotFoundError,
    EmptyContentError,
    InvalidDocumentError,
    ProcessingError,
)

_MAX_MEMBERS = 4096
_MAX_EXPANDED_BYTES = 256 * 1024 * 1024
_IMAGE_EXTENSIONS = {
    "image/png": ".png",
    "image/jpeg": ".jpg",
    "image/gif": ".gif",
    "image/webp": ".webp",
    "image/bmp": ".bmp",
    "image/tiff": ".tiff",
}
_MATH_NS = "http://schemas.openxmlformats.org/officeDocument/2006/math"
_MATH_CONTAINERS = {
    "oMath",
    "oMathPara",
    "r",
    "e",
    "num",
    "den",
    "sub",
    "sup",
    "deg",
    "fName",
    "lim",
    "box",
    "borderBox",
}
_MATH_STRUCTURES = {
    "t",
    "f",
    "sSup",
    "sSub",
    "sSubSup",
    "rad",
    "nary",
    "d",
    "func",
    "limLow",
    "limUpp",
    "acc",
    "bar",
    "groupChr",
    "m",
    "mr",
    "eqArr",
}


def _unsupported_math(element):
    unknown = set()

    def inspect(node):
        name = etree.QName(node)
        if name.namespace == _MATH_NS:
            if name.localname.endswith("Pr"):
                return
            if name.localname not in _MATH_CONTAINERS | _MATH_STRUCTURES:
                unknown.add(name.localname)
        for item in node:
            inspect(item)

    inspect(element)
    return sorted(unknown)


def _math_text(node):
    name = etree.QName(node).localname

    def child(tag):
        item = node.find("{" + _MATH_NS + "}" + tag)
        return _math_text(item) if item is not None else ""

    def property_value(container, tag, default):
        item = node.find("{" + _MATH_NS + "}" + container + "/{" + _MATH_NS + "}" + tag)
        return item.get("{" + _MATH_NS + "}val", default) if item is not None else default

    if name == "t":
        return node.text or ""
    if name.endswith("Pr"):
        return ""
    if name == "f":
        if property_value("fPr", "type", "bar") == "noBar":
            return "stack(" + child("num") + "; " + child("den") + ")"
        return "(" + child("num") + ")/(" + child("den") + ")"
    if name in {"sSup", "sSub", "sSubSup"}:
        base = child("e")
        if re.fullmatch(r"\w+", base) is None:
            base = "(" + base + ")"
        return (
            base
            + ("_(" + child("sub") + ")" if name != "sSup" else "")
            + ("^(" + child("sup") + ")" if name != "sSub" else "")
        )
    if name == "rad":
        degree = child("deg")
        return (
            ("root(" + degree + ", " + child("e") + ")") if degree else "sqrt(" + child("e") + ")"
        )
    if name == "nary":
        symbol = property_value("naryPr", "chr", "∫")
        return (
            symbol
            + ("_(" + child("sub") + ")" if child("sub") else "")
            + ("^(" + child("sup") + ")" if child("sup") else "")
            + " ("
            + child("e")
            + ")"
        )
    if name == "d":
        entries = [_math_text(item) for item in node.findall("{" + _MATH_NS + "}e")]
        return (
            property_value("dPr", "begChr", "(")
            + property_value("dPr", "sepChr", "|").join(entries)
            + property_value("dPr", "endChr", ")")
        )
    if name == "func":
        return child("fName") + "(" + child("e") + ")"
    if name in {"limLow", "limUpp"}:
        return child("e") + ("_(" if name == "limLow" else "^(") + child("lim") + ")"
    if name in {"acc", "bar", "groupChr"}:
        return property_value(name + "Pr", "chr", name) + "(" + child("e") + ")"
    if name in {"m", "mr", "eqArr"}:
        separator = "; " if name in {"m", "eqArr"} else ", "
        return (
            "["
            + separator.join(
                _math_text(item) for item in node if not etree.QName(item).localname.endswith("Pr")
            )
            + "]"
        )
    text = "".join(_math_text(item) for item in node)
    return text if name in _MATH_CONTAINERS else "[unsupported OMML " + name + ": " + text + "]"


def _paragraph_text(paragraph):
    def ordered(node):
        if node.tag in {"{" + _MATH_NS + "}oMath", "{" + _MATH_NS + "}oMathPara"}:
            return _math_text(node)
        if node.tag == qn("w:t"):
            return node.text or ""
        if node.tag == qn("w:tab"):
            return "\t"
        if node.tag in {qn("w:br"), qn("w:cr")}:
            return "\n"
        if node.tag in {qn("w:drawing"), qn("w:pict"), qn("w:txbxContent")}:
            return ""
        return "".join(ordered(item) for item in node)

    from tools.ocr_support import normalize_ocr_text

    return normalize_ocr_text(ordered(paragraph._p))


def _cell_text(cell):
    parts = []
    for element in cell._tc.iterchildren():
        if element.tag == qn("w:p"):
            text = _paragraph_text(Paragraph(element, cell)).strip()
        elif element.tag == qn("w:tbl"):
            text = _table_text(Table(element, cell))
        else:
            continue
        if text:
            parts.append(text)
    return " / ".join(parts).replace("\t", " ")


def _open_document(path: Path):
    if not path.is_file():
        raise AssetNotFoundError("The Word document was not found.")
    try:
        with ZipFile(path) as archive:
            members = archive.infolist()
            if (
                len(members) > _MAX_MEMBERS
                or sum(member.file_size for member in members) > _MAX_EXPANDED_BYTES
                or any(member.flag_bits & 1 for member in members)
                or "word/document.xml" not in archive.namelist()
            ):
                raise InvalidDocumentError(
                    "The Word document exceeds safe package limits or is encrypted."
                )
        return WordDocument(path)
    except InvalidDocumentError:
        raise
    except (BadZipFile, OSError, ValueError, KeyError) as exc:
        raise InvalidDocumentError("The file is not a readable Word document.") from exc
    except Exception as exc:
        raise InvalidDocumentError("The Word document could not be read.") from exc


def _blocks(document):
    for position, element in enumerate(document.element.body.iterchildren(), 1):
        if element.tag == qn("w:p"):
            yield f"body-{position:06d}", Paragraph(element, document)
        elif element.tag == qn("w:tbl"):
            yield f"body-{position:06d}", Table(element, document)


def _heading_level(paragraph: Paragraph) -> int | None:
    style, seen = paragraph.style, set()
    properties = paragraph._p.pPr
    while True:
        level = properties.find(qn("w:outlineLvl")) if properties is not None else None
        if level is not None:
            value = level.get(qn("w:val"))
            if value is not None and value.isdecimal() and 0 <= int(value) <= 9:
                return int(value) + 1 if int(value) < 9 else None
        if style is None or style.style_id in seen:
            break
        seen.add(style.style_id)
        match = re.fullmatch(r"heading\s*([1-9])", style.name or "", re.IGNORECASE)
        if match:
            return int(match.group(1))
        properties = style.element.pPr
        style = style.base_style
    return None


def _table_text(table: Table) -> str:
    rows, seen = [], set()
    for row in table.rows:
        values = []
        for cell in row.cells:
            if cell._tc in seen:
                values.append("")
                continue
            seen.add(cell._tc)
            values.append(_cell_text(cell))
        if any(values):
            rows.append("\t".join(values))
    return "\n".join(rows).strip()


def _document_metadata(document) -> dict:
    properties = document.core_properties
    metadata = {}
    for name in ("title", "author", "subject", "keywords", "category", "language"):
        value = getattr(properties, name, None)
        if isinstance(value, str) and value.strip():
            metadata[f"document_{name}"] = value.strip()[:512]
    for name in ("created", "modified"):
        value = getattr(properties, name, None)
        if value is not None:
            metadata[f"document_{name}"] = value.isoformat()
    return metadata


def _image_relationships(element):
    identifiers = []
    for node in element.iter():
        if node.tag == qn("a:blip"):
            identifier = node.get(qn("r:embed"))
        elif node.tag.endswith("}imagedata"):
            identifier = node.get(qn("r:id"))
        else:
            continue
        if identifier and identifier not in identifiers:
            identifiers.append(identifier)
    return identifiers


def _extract_images(document, path, image_dir, max_images, max_image_bytes):
    if type(max_images) is not int or not 0 <= max_images <= 128:
        raise ValueError("max_images must be between zero and 128.")
    if type(max_image_bytes) is not int or not 1 <= max_image_bytes <= 64 * 1024 * 1024:
        raise ValueError("max_image_bytes must be between one byte and 64 MiB.")
    destination = Path(image_dir) if image_dir is not None else path.parent / "images"
    # References stay relative to the asset directory even when it is staged.
    try:
        relative_dir = destination.resolve().relative_to(path.parent.resolve())
    except ValueError as exc:
        raise ValueError("The image directory must be inside the source asset directory.") from exc
    records, written = [], set()
    for element_id, block in _blocks(document):
        for identifier in _image_relationships(block._element):
            relationship = document.part.rels.get(identifier)
            if relationship is None or relationship.is_external:
                continue
            part = relationship.target_part
            mime = part.content_type
            if not mime.startswith("image/"):
                continue
            blob = part.blob
            digest = hashlib.sha256(blob).hexdigest()
            record = {
                "element_id": element_id,
                "image_id": digest,
                "content_type": mime,
                "byte_count": len(blob),
            }
            extension = _IMAGE_EXTENSIONS.get(mime)
            if not extension:
                record["status"] = "unsupported_format"
            elif len(blob) > max_image_bytes:
                record["status"] = "too_large"
            elif digest not in written and len(written) >= max_images:
                record["status"] = "limit_reached"
            else:
                filename = digest + extension
                if digest not in written:
                    destination.mkdir(parents=True, exist_ok=True)
                    (destination / filename).write_bytes(blob)
                written.add(digest)
                record.update(status="available", reference=(relative_dir / filename).as_posix())
            records.append(record)
    return records


def extract_docx_images(
    docx_path: str,
    *,
    image_dir: str | Path | None = None,
    max_images: int = 12,
    max_image_bytes: int = 10 * 1024 * 1024,
) -> list[dict]:
    path = Path(docx_path)
    document = _open_document(path)
    try:
        return _extract_images(document, path, image_dir, max_images, max_image_bytes)
    except OSError as exc:
        raise ProcessingError("The embedded Word images could not be stored.") from exc


def extract_docx_documents(
    docx_path: str,
    asset_id: str | None = None,
    source_name: str | None = None,
    *,
    image_dir: str | Path | None = None,
    max_images: int = 12,
    max_image_bytes: int = 10 * 1024 * 1024,
) -> list[Document]:
    path = Path(docx_path)
    document = _open_document(path)
    source = Path(source_name or path.name).name
    try:
        images = _extract_images(document, path, image_dir, max_images, max_image_bytes)
        common = {
            "source": source,
            "source_type": "docx",
            "modality": "docx",
            **_document_metadata(document),
        }
        if asset_id is not None:
            common["asset_id"] = asset_id
        documents, headings, images_by_element = [], {}, {}
        for image in images:
            images_by_element.setdefault(image["element_id"], []).append(image)
        paragraph_index = table_index = 0
        for element_id, block in _blocks(document):
            metadata = {**common, "element_id": element_id, "record_index": len(documents)}
            if isinstance(block, Paragraph):
                paragraph_index += 1
                text = _paragraph_text(block).strip()
                level = _heading_level(block)
                metadata.update(
                    content_type="heading" if level else "paragraph", paragraph=paragraph_index
                )
                if level and text:
                    headings = {depth: title for depth, title in headings.items() if depth < level}
                    headings[level] = text
                    metadata["heading_level"] = level
            else:
                table_index += 1
                text = _table_text(block)
                metadata.update(
                    content_type="table",
                    table_index=table_index,
                    table_data=json.dumps(
                        [[_cell_text(cell) for cell in row.cells] for row in block.rows],
                        ensure_ascii=False,
                    ),
                )
            if any(True for _ in block._element.iter("{" + _MATH_NS + "}oMath")):
                metadata.update(
                    extraction_method="ooxml_math",
                    uncertainty="Native Word equations were linearized from OMML; visual layout was not rendered.",
                )
                unsupported = _unsupported_math(block._element)
                if unsupported:
                    metadata["uncertainty"] += (
                        " Unsupported OMML structures: "
                        + ", ".join(unsupported)
                        + "; linearization may be incomplete."
                    )
            path_names = [title for _, title in sorted(headings.items())]
            if path_names:
                metadata.update(
                    section=path_names[-1], section_path=json.dumps(path_names, ensure_ascii=False)
                )
            owned_images = images_by_element.get(element_id, [])
            if owned_images:
                metadata.update(
                    image_count=len(owned_images),
                    image_references=json.dumps(owned_images, ensure_ascii=False, sort_keys=True),
                )
            if text:
                documents.append(Document(page_content=text, metadata=metadata))
    except (OSError, ValueError, TypeError) as exc:
        raise InvalidDocumentError("The Word document content could not be extracted.") from exc
    if not documents:
        raise EmptyContentError("The Word document contains no readable paragraphs or tables.")
    return documents


_CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_DRAWING_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
_DIAGRAM_NS = "http://schemas.openxmlformats.org/drawingml/2006/diagram"


def _part_xml(document, identifier):
    relationship = document.part.rels.get(identifier)
    if relationship is None or relationship.is_external:
        return None
    parser = etree.XMLParser(resolve_entities=False, no_network=True, recover=False)
    return etree.fromstring(relationship.target_part.blob, parser)


def _chart_data(root):
    namespace = {"c": _CHART_NS, "a": _DRAWING_NS}
    title = " ".join(root.xpath(".//c:title//a:t/text()", namespaces=namespace))
    chart_types = [
        etree.QName(node).localname
        for node in root.iter()
        if etree.QName(node).namespace == _CHART_NS
        and etree.QName(node).localname.endswith("Chart")
    ]
    lines, series_data = ["Native chart" + (": " + title if title else "")], []
    for number, series in enumerate(root.xpath(".//c:ser", namespaces=namespace), 1):
        names = series.xpath("./c:tx//c:v/text() | ./c:tx//a:t/text()", namespaces=namespace)
        name = " ".join(names) or f"Series {number}"
        categories = {
            point.get("idx"): " ".join(point.xpath("./c:v/text()", namespaces=namespace))
            for point in series.xpath("./c:cat//c:pt | ./c:xVal//c:pt", namespaces=namespace)
        }
        values = {
            point.get("idx"): " ".join(point.xpath("./c:v/text()", namespaces=namespace))
            for point in series.xpath("./c:val//c:pt | ./c:yVal//c:pt", namespaces=namespace)
        }
        points = [
            {"label": categories.get(index, f"Point {index}"), "value": value}
            for index, value in sorted(
                values.items(), key=lambda item: int(item[0]) if (item[0] or "").isdigit() else 0
            )
        ]
        lines.append(name + ":")
        lines.extend(f"{point['label']}: {point['value']}" for point in points)
        series_data.append({"name": name, "points": points})
    axes = [
        " ".join(axis.xpath(".//c:title//a:t/text()", namespaces=namespace))
        for axis in root.xpath(".//c:catAx | .//c:valAx", namespaces=namespace)
    ]
    if any(axes):
        lines.append("Axis titles: " + "; ".join(axis for axis in axes if axis))
    return "\n".join(lines), {
        "chart_type": chart_types,
        "title": title,
        "axes": axes,
        "series": series_data,
    }


def _native_structure(data, warnings):
    encoded = json.dumps(data, ensure_ascii=False)
    if len(encoded) <= 64000:
        return encoded
    warnings.append(
        "Native visual structured metadata exceeded the supported limit; its extracted text was retained."
    )
    return json.dumps({"status": "metadata_limit", "original_characters": len(encoded)})


def _native_objects(document, common):
    records, warnings = [], []
    for element_id, block in _blocks(document):
        base = {
            **common,
            "element_id": element_id,
            "extraction_method": "ooxml",
            "processing_status": "completed",
        }
        for node in block._element.iter():
            if node.tag == "{" + _CHART_NS + "}chart":
                try:
                    root = _part_xml(document, node.get(qn("r:id")))
                    if root is None:
                        warnings.append("A native Word chart relationship could not be read.")
                        continue
                    text, data = _chart_data(root)
                    if not any(series["points"] for series in data["series"]):
                        warnings.append(
                            "A native Word chart has no cached values; its external workbook or rendering is unavailable."
                        )
                    records.append(
                        Document(
                            page_content=text,
                            metadata={
                                **base,
                                "content_type": "chart_data",
                                "visual_artifact_type": "chart",
                                "visual_structure": _native_structure(data, warnings),
                            },
                        )
                    )
                    warnings.append(
                        "Native Word chart values and axis titles were extracted; the original chart layout was not rendered for VLM analysis."
                    )
                except (etree.XMLSyntaxError, ValueError, KeyError):
                    warnings.append(
                        "A native Word chart could not be parsed; its visual layout was not rendered."
                    )
            elif node.tag == "{" + _DIAGRAM_NS + "}relIds":
                try:
                    root = _part_xml(document, node.get(qn("r:dm")))
                    if root is None:
                        warnings.append("A SmartArt data relationship could not be read.")
                        continue
                    points, links = {}, []
                    for point in root.iter("{" + _DIAGRAM_NS + "}pt"):
                        labels = [
                            part.text for part in point.iter("{" + _DRAWING_NS + "}t") if part.text
                        ]
                        if labels:
                            points[point.get("modelId", str(len(points)))] = " ".join(labels)
                    for connection in root.iter("{" + _DIAGRAM_NS + "}cxn"):
                        start, end = connection.get("srcId"), connection.get("destId")
                        if start in points and end in points:
                            links.append(
                                {
                                    "from": points[start],
                                    "to": points[end],
                                    "type": connection.get("type", "connection"),
                                }
                            )
                    if points:
                        text = "SmartArt elements:\n" + "\n".join(points.values())
                        if links:
                            text += "\nConnections:\n" + "\n".join(
                                f"{link['from']} -> {link['to']} ({link['type']})" for link in links
                            )
                        records.append(
                            Document(
                                page_content=text[:64000],
                                metadata={
                                    **base,
                                    "content_type": "diagram_data",
                                    "visual_artifact_type": "smartart",
                                    "visual_structure": _native_structure(
                                        {"elements": points, "connections": links}, warnings
                                    ),
                                },
                            )
                        )
                    warnings.append(
                        "SmartArt native labels and connections were extracted; its visual layout requires a Word renderer and was not rendered."
                    )
                except (etree.XMLSyntaxError, ValueError, KeyError):
                    warnings.append(
                        "A SmartArt object could not be parsed; its visual layout was not rendered."
                    )
        shape_labels = [
            node.text for node in block._element.iter("{" + _DRAWING_NS + "}t") if node.text
        ]
        textbox_labels = []
        for textbox in block._element.iter(qn("w:txbxContent")):
            textbox_labels.extend(node.text for node in textbox.iter(qn("w:t")) if node.text)
        if shape_labels or textbox_labels:
            labels = list(dict.fromkeys(shape_labels + textbox_labels))
            records.append(
                Document(
                    page_content="DrawingML labels: " + " / ".join(labels),
                    metadata={
                        **base,
                        "content_type": "diagram_data",
                        "visual_artifact_type": "drawingml",
                    },
                )
            )
            warnings.append(
                "Native DrawingML/text-box labels were extracted; unsupported shape rendering was not sent to the VLM."
            )
        shape_kinds = set()
        for node in block._element.iter():
            name = etree.QName(node)
            if name.namespace == "urn:schemas-microsoft-com:vml" and name.localname in {
                "shape",
                "group",
                "rect",
                "oval",
                "line",
                "polyline",
                "roundrect",
                "arc",
                "curve",
            }:
                shape_kinds.add("VML")
            elif name.namespace in {
                _DRAWING_NS,
                "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingShape",
                "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingGroup",
            } and name.localname in {"sp", "wsp", "wgp", "grpSp", "cxnSp"}:
                shape_kinds.add("DrawingML")
        for kind in sorted(shape_kinds):
            warnings.append(
                f"{kind} shape geometry/layout at {element_id} was not extracted; textless shapes may contain missing information. Optional LibreOffice rendering is required and only runs when visual understanding is enabled."
            )
    return records, list(dict.fromkeys(warnings))


def extract_docx_hybrid(
    source, *, asset_id, source_name, settings, vlm_mode, deadline, progress=None
):
    from tools.document_visuals import (
        check_deadline,
        finish_records,
        merge_metadata,
        notify,
        setting,
        visual_metadata,
    )

    if vlm_mode not in {"disabled", "automatic", "enhanced"}:
        raise ValueError("Unsupported VLM processing mode.")
    path = Path(source)
    check_deadline(deadline)
    notify(progress, "extracting_document_text", source_type="docx")
    limit = min(128, max(1, int(setting(settings, "VLM_MAX_VISUALS", 32))))
    result = visual_metadata(vlm_mode)
    try:
        documents = extract_docx_documents(
            str(path), asset_id=asset_id, source_name=source_name, max_images=limit
        )
    except EmptyContentError:
        documents = []
    document = _open_document(path)
    common = {
        "asset_id": asset_id,
        "source": Path(source_name or path.name).name,
        "source_type": "docx",
        "modality": "docx",
        **_document_metadata(document),
    }
    native_documents, native_warnings = _native_objects(document, common)
    documents.extend(native_documents)
    result["warnings"].extend(native_warnings)
    images = _extract_images(document, path, None, limit, 10 * 1024 * 1024)
    for image in images:
        if image.get("reference"):
            result["files"].append(image["reference"])
        elif image["status"] != "available":
            result["warnings"].append(f"An embedded Word image was omitted ({image['status']}).")
    by_element = {}
    for record in documents:
        record.metadata.setdefault("extraction_method", "python_docx")
        if "Unsupported OMML" in record.metadata.get("uncertainty", ""):
            result["warnings"].append(
                record.metadata["element_id"] + ": " + record.metadata["uncertainty"]
            )
        if record.metadata.get("content_type") == "table":
            record.metadata.setdefault(
                "table_data",
                json.dumps(
                    [line.split("\t") for line in record.page_content.splitlines()],
                    ensure_ascii=False,
                ),
            )
        by_element.setdefault(record.metadata["element_id"], []).append(record)
    blocks = list(_blocks(document))
    element_context, headings = {}, {}
    for index, (element_id, block) in enumerate(blocks):
        if isinstance(block, Paragraph):
            level = _heading_level(block)
            if level and _paragraph_text(block).strip():
                headings = {depth: title for depth, title in headings.items() if depth < level}
                headings[level] = _paragraph_text(block).strip()
        context = {}
        names = [title for _, title in sorted(headings.items())]
        if names:
            context.update(section=names[-1], section_path=json.dumps(names, ensure_ascii=False))
        neighbors = []
        for _, neighbor in blocks[max(0, index - 1) : index + 2]:
            neighbors.append(
                _paragraph_text(neighbor)
                if isinstance(neighbor, Paragraph)
                else _table_text(neighbor)
            )
        nearby = "\n".join(text for text in neighbors if text.strip())[:2000]
        caption = " ".join(
            text.strip()
            for text in neighbors
            if re.match(r"^(fig(?:ure)?\b|diagram\b|chart\b|flowchart\b)", text.strip(), re.I)
        )[:512]
        element_context[element_id] = {**context, "nearby_text": nearby, "caption": caption}
        for record in by_element.get(element_id, []):
            for key, value in context.items():
                record.metadata.setdefault(key, value)
    notify(
        progress,
        "text_extraction_completed",
        records=[
            record.model_dump()
            for record in documents
            if record.metadata.get("content_type") != "table"
        ],
    )
    notify(
        progress,
        "table_extraction_completed",
        records=[
            record.model_dump()
            for record in documents
            if record.metadata.get("content_type") == "table"
        ],
    )
    artifacts = []
    if (
        vlm_mode != "disabled"
        and (native_documents or native_warnings)
        and setting(settings, "DOCX_RENDER_ENABLED", True)
    ):
        from tools.docx_rendering import render_word, rendered_candidates

        notify(progress, "rendering_word_visuals")
        rendered, warning = render_word(path, settings=settings, deadline=deadline)
        if warning:
            result["warnings"].append(warning)
        elif rendered:
            artifacts.extend(
                rendered_candidates(path, rendered, settings=settings, deadline=deadline)
            )
            result["files"].append("rendered/" + rendered.name)
            result["warnings"].append(
                "Complex Word visuals were rendered with LibreOffice; rendered page locators supplement native element references."
            )
    if vlm_mode != "disabled" and images:
        notify(progress, "detecting_visual_content", source_type="docx")
        for image in images:
            check_deadline(deadline)
            if image["status"] != "available":
                artifacts.append(
                    {
                        "reference": image.get("reference", ""),
                        "artifact_type": "image",
                        "element_id": image["element_id"],
                        "processing_status": "skipped_budget"
                        if image["status"]
                        in {"limit", "limit_reached", "too_large", "omitted", "skipped"}
                        else "failed",
                        "reason": image["status"],
                    }
                )
                continue
            context = element_context.get(image["element_id"], {})
            artifact = {
                "reference": image["reference"],
                "artifact_type": "image",
                "element_id": image["element_id"],
                **context,
                "priority": 0.9 if context.get("caption") else 0.5,
            }
            if context.get("caption"):
                artifact["meaningful"] = True
            artifacts.append(artifact)
    if artifacts:
        from tools.visual_understanding import analyze_artifacts

        visual_docs, visual_result = analyze_artifacts(
            path,
            artifacts,
            asset_id=asset_id,
            source_name=source_name,
            source_type="docx",
            vlm_mode=vlm_mode,
            settings=settings,
            deadline=deadline,
            progress=progress,
        )
        for record in visual_docs:
            context = element_context.get(record.metadata.get("element_id"), {})
            for key in ("section", "section_path"):
                if context.get(key):
                    record.metadata.setdefault(key, context[key])
        documents.extend(visual_docs)
        merge_metadata(result, visual_result)
        if visual_docs:
            notify(progress, "merging_multimodal", source_type="docx")
    result["detected_visual_count"] = max(result["detected_visual_count"], len(artifacts))
    result["files"] = sorted(set(result["files"]))
    result["warnings"] = list(dict.fromkeys(result["warnings"]))
    documents.sort(
        key=lambda record: (
            record.metadata["element_id"],
            int(record.metadata.get("content_type") == "visual_description"),
        )
    )
    if not documents:
        raise EmptyContentError(
            "The Word document contains no readable text, tables, native chart data, or analyzed visual evidence. "
            + " ".join(result["warnings"])
        )
    return finish_records(documents, vlm_mode=vlm_mode), result
